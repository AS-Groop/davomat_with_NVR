# -*- coding: utf-8 -*-
"""
yoqlama_doc.py — Yo'qlama varaqasini Word (.docx) va PDF qilib yasaydi.

Ishlatilishi (mobile_admin.py dan):
    import yoqlama_doc
    p = yoqlama_doc.build_docx(item)   # -> Path
    p = yoqlama_doc.build_pdf(item)    # -> Path

item tuzilishi (yoqlama.json dagi bitta yozuv):
    {id, title, date, time, place, employees:[{card_number, full_name}], marks:{card:bool}}
"""

import os
import re
from pathlib import Path
from datetime import datetime

BASE_DIR = Path(".")
OUT_DIR = BASE_DIR / "yoqlamalar"
DEPT_FILE = BASE_DIR / "employee_departments.json"

try:
    import tashkilot as _tashkilot
    _ORG_FULL = _tashkilot.get().get("full", "")
except Exception:
    _ORG_FULL = ""
ORG_NAME = os.getenv("ORG_FULL_NAME", _ORG_FULL)

UZ_MONTHS = ["yanvar", "fevral", "mart", "aprel", "may", "iyun",
             "iyul", "avgust", "sentabr", "oktabr", "noyabr", "dekabr"]


def _fmt_date(iso):
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", str(iso or ""))
    if not m:
        d = datetime.now()
        return f"{d.year}-yil {d.day}-{UZ_MONTHS[d.month - 1]}"
    y, mo, da = int(m.group(1)), int(m.group(2)), int(m.group(3))
    return f"{y}-yil {da}-{UZ_MONTHS[mo - 1]}"


def _dept_pos(card):
    """employee_departments.json dan (guruh, lavozim)."""
    try:
        import json
        d = json.loads(DEPT_FILE.read_text(encoding="utf-8"))
        rec = d.get(str(card)) or {}
        dept = str(rec.get("department", "") or "").strip()
        pos = str(rec.get("position", "") or "").strip()
        if dept and pos:
            return f"{dept}, {pos}"
        return dept or pos or ""
    except Exception:
        return ""


def _fmt_name(full):
    """'ALIYEV SARDOR BAXTIYOR O‘G‘LI' -> 'Aliyev Sardor Baxtiyor o‘g‘li'."""
    s = str(full or "").strip()
    if not s:
        return ""
    keep = {"o‘g‘li", "o'g'li", "qizi", "ugli"}
    out = []
    for w in s.split():
        lw = w.lower()
        out.append(lw if lw in keep else (lw[:1].upper() + lw[1:]))
    return " ".join(out)


def _rows(item):
    """[(n, fio, dept_pos, '+'/'-'), ...]"""
    marks = item.get("marks", {}) or {}
    emps = item.get("employees", []) or []
    out = []
    for i, e in enumerate(emps, 1):
        card = str(e.get("card_number", ""))
        out.append((i, _fmt_name(e.get("full_name", "")), _dept_pos(card),
                    "+" if marks.get(card) else "−"))
    return out


def _counts(item):
    marks = item.get("marks", {}) or {}
    yes = sum(1 for v in marks.values() if v)
    return yes, len(marks)


def _safe(s):
    return re.sub(r"[^\w\-]+", "_", str(s or ""), flags=re.UNICODE)[:50] or "yoqlama"


def _header_line(item):
    parts = [_fmt_date(item.get("date"))]
    if item.get("time"):
        parts.append(f"soat {item['time']}")
    if item.get("place"):
        parts.append(str(item["place"]))
    return " · ".join(parts)


SIG_DIR = BASE_DIR / "explanations" / "signatures"


def _conductor(item):
    """Yo'qlamani o'tkazgan xodim: (F.I.Sh., imzo fayli yo'li yoki None)."""
    c = item.get("conductor") or {}
    name = _fmt_name(c.get("full_name", "")) or ""
    sig = str(item.get("signature_file", "") or "")
    p = SIG_DIR / sig if sig else None
    return name, (str(p) if p and p.exists() else None)


# ─────────────────────────── WORD ───────────────────────────

def build_docx(item):
    """Yo'qlama varaqasini .docx qilib yasaydi. Path qaytaradi."""
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUT_DIR / f"{item.get('id','yoq')}_{_safe(item.get('title'))}.docx"

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2); sec.bottom_margin = Cm(2)
    sec.left_margin = Cm(2.5); sec.right_margin = Cm(1.5)

    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)

    def para(text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=4):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(after)
        r = p.add_run(text)
        r.font.name = "Times New Roman"; r.font.size = Pt(size); r.bold = bold
        return p

    para(ORG_NAME, size=11, after=10)
    para("YO‘QLAMA VARAQASI", size=14, bold=True, after=6)
    para(str(item.get("title", "")), size=13, bold=True, after=2)
    para(_header_line(item), size=11, after=12)

    rows = _rows(item)
    tab = doc.add_table(rows=1, cols=4)
    tab.style = "Table Grid"
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tab.rows[0].cells
    for i, txt in enumerate(["№", "F.I.Sh.", "Guruh / lavozim", "Qatnashdi"]):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(txt); r.bold = True
        r.font.name = "Times New Roman"; r.font.size = Pt(11)

    for n, fio, dp, mark in rows:
        c = tab.add_row().cells
        vals = [str(n), fio, dp, mark]
        for i, v in enumerate(vals):
            c[i].text = ""
            p = c[i].paragraphs[0]
            p.alignment = (WD_ALIGN_PARAGRAPH.CENTER if i in (0, 3)
                           else WD_ALIGN_PARAGRAPH.LEFT)
            r = p.add_run(v)
            r.font.name = "Times New Roman"; r.font.size = Pt(11)

    for i, w in enumerate([Cm(1.2), Cm(6.5), Cm(7.0), Cm(2.3)]):
        for row in tab.rows:
            row.cells[i].width = w

    yes, total = _counts(item)
    doc.add_paragraph()
    para(f"Jami qatnashdi: {yes} / {total}", size=12, bold=True,
         align=WD_ALIGN_PARAGRAPH.RIGHT, after=18)

    # ── Yo'qlamani o'tkazgan xodim va uning imzosi ──
    cname, csig = _conductor(item)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    r = p.add_run("Yo‘qlamani o‘tkazdi:")
    r.font.name = "Times New Roman"; r.font.size = Pt(12)

    ft = doc.add_table(rows=1, cols=2)
    ft.alignment = WD_TABLE_ALIGNMENT.LEFT
    left, right = ft.rows[0].cells
    left.width = Cm(8.5); right.width = Cm(8.0)

    pl = left.paragraphs[0]
    rl = pl.add_run(cname or "_____________________")
    rl.font.name = "Times New Roman"; rl.font.size = Pt(12); rl.bold = bool(cname)

    pr = right.paragraphs[0]
    pr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if csig:
        try:
            pr.add_run().add_picture(csig, width=Cm(4.2))
        except Exception:
            rr = pr.add_run("_____________________")
            rr.font.name = "Times New Roman"; rr.font.size = Pt(12)
    else:
        rr = pr.add_run("Imzo: _____________________")
        rr.font.name = "Times New Roman"; rr.font.size = Pt(12)

    doc.save(str(out))
    return out


# ─────────────────────────── PDF ───────────────────────────

def build_pdf(item):
    """Yo'qlama varaqasini .pdf qilib yasaydi. Path qaytaradi."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer,
                                    Table, TableStyle)

    font, font_b = "Helvetica", "Helvetica-Bold"
    try:
        import explanation_pdf as ep
        f = ep._ensure_fonts()          # ("ExplReg", "ExplBold") yoki False
        if f:
            font, font_b = f
    except Exception:
        pass

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUT_DIR / f"{item.get('id','yoq')}_{_safe(item.get('title'))}.pdf"

    doc = SimpleDocTemplate(str(out), pagesize=A4,
                            topMargin=18 * mm, bottomMargin=18 * mm,
                            leftMargin=22 * mm, rightMargin=15 * mm)
    S = lambda n, sz, b=False, al=1, sp=4: ParagraphStyle(
        n, fontName=(font_b if b else font), fontSize=sz, leading=sz * 1.35,
        alignment=al, spaceAfter=sp)

    story = [
        Paragraph(ORG_NAME, S("org", 10.5, sp=10)),
        Paragraph("YO‘QLAMA VARAQASI", S("t", 14, True, sp=6)),
        Paragraph(str(item.get("title", "")), S("n", 12.5, True, sp=2)),
        Paragraph(_header_line(item), S("d", 10.5, sp=12)),
    ]

    data = [["№", "F.I.Sh.", "Guruh / lavozim", "Qatnashdi"]]
    cell = ParagraphStyle("c", fontName=font, fontSize=9.5, leading=12, alignment=0)
    for n, fio, dp, mark in _rows(item):
        data.append([str(n), Paragraph(fio, cell), Paragraph(dp, cell), mark])

    tab = Table(data, colWidths=[12 * mm, 62 * mm, 68 * mm, 22 * mm], repeatRows=1)
    tab.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, 0), font_b),
        ("FONTSIZE", (0, 0), (-1, 0), 9.5),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF4F8")),
        ("FONTNAME", (0, 1), (-1, -1), font),
        ("FONTSIZE", (0, 1), (-1, -1), 9.5),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#8AA0AE")),
        ("ALIGN", (0, 0), (0, -1), "CENTER"),
        ("ALIGN", (3, 0), (3, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(tab)

    yes, total = _counts(item)
    story += [
        Spacer(1, 8 * mm),
        Paragraph(f"Jami qatnashdi: {yes} / {total}", S("s", 12, True, al=2, sp=16)),
        Spacer(1, 10 * mm),
    ]

    # ── Yo'qlamani o'tkazgan xodim va uning imzosi ──
    from reportlab.platypus import Image as RLImage
    cname, csig = _conductor(item)
    story.append(Paragraph("Yo‘qlamani o‘tkazdi:", S("cl", 11, al=0, sp=6)))

    sig_cell = ""
    if csig:
        try:
            sig_cell = RLImage(csig, width=42 * mm, height=16 * mm, kind="proportional")
        except Exception:
            sig_cell = Paragraph("_____________________", S("x", 11, al=1))
    else:
        sig_cell = Paragraph("Imzo: _____________________", S("x", 11, al=1))

    name_cell = Paragraph(cname or "_____________________",
                          S("cn", 11.5, bool(cname), al=0))
    ft = Table([[name_cell, sig_cell]], colWidths=[92 * mm, 72 * mm])
    ft.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "BOTTOM"),
        ("ALIGN", (1, 0), (1, 0), "CENTER"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    story.append(ft)

    doc.build(story)
    return out

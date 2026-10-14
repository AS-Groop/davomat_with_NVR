# -*- coding: utf-8 -*-
"""reports_xlsx.py — HR/super_admin uchun Excel hisobotlar.

Ikki turi:
  build_rating_xlsx(ym)  — oylik reyting (ish soatlari, kech qolish,
                            kelmagan kunlar) — BARCHA xodimlar bo'yicha
                            (faqat "top_work"dagi emas — 0 kunliklar ham).
  build_journal_xlsx(ym) — kirish-chiqish jurnali: har xodim, har kun,
                            kirish/chiqish vaqti va jami soat.

Ikkalasi ham `exports/` papkasiga yoziladi (.gitignore'da — haqiqiy
davomat ma'lumoti bo'lgani uchun git'ga ketmaydi).
"""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

BASE_DIR = Path(".")
EXPORTS_DIR = BASE_DIR / "exports"
USERS_FILE = BASE_DIR / "users.xlsx"


def _employees_meta() -> dict:
    """users.xlsx dan {card: {full_name, department, position}} (tartib saqlanadi)."""
    import openpyxl

    wb = openpyxl.load_workbook(USERS_FILE, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return {}
    hdr = [str(c or "").strip().lower() for c in rows[0]]

    def idx(*names):
        for n in names:
            if n in hdr:
                return hdr.index(n)
        return None

    ci = idx("card_number")
    fi, ii, shi = idx("familiya"), idx("ism"), idx("sharif")
    di = idx("bo'lim", "bolim")
    pi = idx("lavozim")
    if ci is None:
        return {}

    out = {}
    for row in rows[1:]:
        card = "".join(ch for ch in str(row[ci] or "") if ch.isdigit())
        if not card:
            continue
        full = " ".join(
            str(row[i]).strip() for i in (fi, ii, shi) if i is not None and row[i]
        ).strip()
        out[card] = {
            "full_name": full or card,
            "department": str(row[di] or "") if di is not None else "",
            "position": str(row[pi] or "") if pi is not None else "",
        }
    return out


def _month_report(ym: str) -> dict:
    cur_path = BASE_DIR / "monthly_report.json"
    try:
        cur = json.loads(cur_path.read_text(encoding="utf-8"))
        if str(cur.get("year_month")) == ym:
            return cur
    except Exception:
        pass
    p = BASE_DIR / "monthly_reports" / f"{ym}.json"
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        return {}


def build_rating_xlsx(ym: str) -> Path:
    """Oylik reyting — BARCHA xodimlar (ish soati bo'yicha kamayish tartibida)."""
    import openpyxl
    from openpyxl.styles import Font

    meta = _employees_meta()
    rep = _month_report(ym)
    work_by_card = {r["card_number"]: r for r in rep.get("top_work", [])}
    late_by_card = {r["card_number"]: r for r in rep.get("top_late", [])}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ym
    headers = ["№", "F.I.Sh.", "Bo'lim", "Lavozim", "Ishlangan kun",
               "Jami soat", "O'rtacha/kun", "Kech qolgan kun", "Kelmagan kun"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)

    rows = []
    for card, m in meta.items():
        w = work_by_card.get(card, {})
        l = late_by_card.get(card, {})
        rows.append((
            w.get("total_minutes", 0),
            m["full_name"], m["department"], m["position"],
            w.get("days_worked", 0), w.get("total_hours_str", "0:00"),
            w.get("avg_hours_str", "0:00"), l.get("late_count", 0),
            w.get("absent_days", 0),
        ))
    rows.sort(key=lambda r: r[0], reverse=True)
    for i, r in enumerate(rows, 1):
        ws.append([i, *r[1:]])

    widths = (5, 30, 20, 22, 12, 10, 12, 12, 12)
    for col, width in zip(ws.columns, widths):
        ws.column_dimensions[col[0].column_letter].width = width

    EXPORTS_DIR.mkdir(exist_ok=True)
    path = EXPORTS_DIR / f"reyting_{ym}.xlsx"
    wb.save(path)
    return path


def build_journal_xlsx(ym: str) -> Path:
    """Kirish-chiqish jurnali — har xodim, har kun (bitta qator)."""
    import openpyxl
    from openpyxl.styles import Font

    import biocontrol_api as api
    from analyze_biocontrol_api import month_range, calc_day_minutes

    meta = _employees_meta()
    start_date, end_date = month_range(ym)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ym
    headers = ["F.I.Sh.", "Bo'lim", "Sana", "Kirish", "Chiqish",
               "Jami soat", "Barcha o'tishlar"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)

    for card, m in meta.items():
        try:
            days = api.get_events_by_day(card, start_date, end_date)
        except Exception:
            continue
        for day in sorted(days.keys()):
            events = sorted(days[day], key=lambda e: e["dt"])
            times = [e["dt"].strftime("%H:%M") for e in events]
            entries = [e for e in events if "вход" in e["type"]]
            exits = [e for e in events if "выход" in e["type"]]
            kirish = entries[0]["dt"].strftime("%H:%M") if entries else ""
            chiqish = exits[-1]["dt"].strftime("%H:%M") if exits else ""
            dm = calc_day_minutes(events)
            jami = f"{dm // 60}:{dm % 60:02d}" if dm else "—"
            ws.append([m["full_name"], m["department"], day, kirish, chiqish,
                       jami, " ".join(times)])

    widths = (30, 20, 12, 9, 9, 10, 40)
    for col, width in zip(ws.columns, widths):
        ws.column_dimensions[col[0].column_letter].width = width
    ws.freeze_panes = "A2"

    EXPORTS_DIR.mkdir(exist_ok=True)
    path = EXPORTS_DIR / f"jurnal_{ym}.xlsx"
    wb.save(path)
    return path


if __name__ == "__main__":
    import sys

    ym = sys.argv[1] if len(sys.argv) > 1 else datetime.now().strftime("%Y-%m")
    print("reyting:", build_rating_xlsx(ym))
    print("jurnal:", build_journal_xlsx(ym))

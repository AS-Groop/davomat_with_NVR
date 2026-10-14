# -*- coding: utf-8 -*-
"""Xodimni tizimdan butunlay chiqarish.

Hisobni o'chirish (accounts.json) — bu faqat ilovaga kirishni to'xtatadi.
Xodim dashboardda, reytingda, kechikkanlar ro'yxatida turaveradi, chunki
u yerlar users.xlsx dan o'qiydi. Ustiga-ustak auto_seed_accounts uni
ko'rib hisobini qayta yaratadi.

Bu modul xodimni HAMMA joydan chiqaradi:
    users.xlsx, accounts.json, tokens, guruh/lavozim, rasm, imtiyozlar,
    ishga kirgan sana, jonli kesh (live_cache.json).

Tegilmaydi:
    monthly_reports/ arxivi va so'rovnoma ovozlari — tarix saqlanadi.

Ishlatish:
    from xodim_chiqarish import remove_employee
    ok, msg, detail = remove_employee("99999999999999")
"""
from __future__ import annotations

import json
import os
import re
import shutil
from datetime import datetime
from pathlib import Path

BASE_DIR = Path(".")

# card -> yozuv ko'rinishidagi oddiy fayllar
DICT_FILES = [
    "employee_departments.json",
    "employee_departments_manual.json",
    "api_photos.json",
    "photo_overrides.json",
    "birthdays_manual.json",
    "hire_dates.json",
    "card_aliases.json",
    "nomzod_chiqarilganlar.json",
    "nomzod_qoshilganlar.json",
    "anketa_ruxsat.json",
    "reasons.json",
    "exit_reasons.json",
    "work_overrides.json",
    "late_excused.json",
    "exit_tracking.json",
]
# {"records": [...]} ko'rinishidagilar
RECORD_FILES = ["exemptions.json"]
# Jonli kesh — dashboard shundan o'qiydi
CACHE_FILE = "live_cache.json"


def _norm(card) -> str:
    return re.sub(r"\D", "", str(card or ""))


def _load(path: Path):
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return None


def _save(path: Path, data) -> None:
    tmp = str(path) + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    os.replace(tmp, path)


def _xlsx_path() -> Path:
    try:
        import live_worker_api as lw
        p = lw.find_existing_employee_file()
        if p:
            return Path(p)
    except Exception:
        pass
    for c in (BASE_DIR / "users.xlsx", BASE_DIR / "employees.xlsx"):
        if c.exists():
            return c
    return BASE_DIR / "users.xlsx"


def find_employee(card: str):
    """Xodim ma'lumoti (users.xlsx bo'yicha) yoki None."""
    try:
        import roles
        return roles.get_record_by_card(_norm(card))
    except Exception:
        return None


def remove_employee(card: str, backup: bool = True) -> tuple[bool, str, dict]:
    """Xodimni butunlay chiqaradi. (ok, xabar, tafsilot) qaytaradi."""
    c = _norm(card)
    if len(c) < 9:
        return False, "Karta raqami noto'g'ri", {}

    rec = find_employee(c)
    name = (rec or {}).get("full_name", "")
    detail: dict = {"card": c, "full_name": name, "cleaned": {}}

    # ── 1) users.xlsx ──
    import openpyxl
    path = _xlsx_path()
    if not path.exists():
        return False, "users.xlsx topilmadi", detail
    try:
        wb = openpyxl.load_workbook(path)
    except PermissionError:
        return False, "users.xlsx Excelda ochiq — yopib qayta urining", detail
    except Exception as e:
        return False, f"users.xlsx o'qilmadi: {e}", detail

    ws = wb.active
    hdr = [re.sub(r"[^a-z]", "", str(x.value or "").lower()) for x in ws[1]]
    col = next((i + 1 for i, h in enumerate(hdr)
                if h in ("cardnumber", "card", "jshshir", "pinfl", "kartaraqami")), None)
    if not col:
        return False, "users.xlsx da karta ustuni topilmadi", detail

    rows = [r for r in range(2, ws.max_row + 1)
            if _norm(ws.cell(r, col).value) == c]
    if not rows:
        return False, f"Bu karta users.xlsx da yo'q ({c})", detail

    if backup:
        bk = BASE_DIR / f"users_backup_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        try:
            shutil.copy2(path, bk)
            detail["backup"] = bk.name
        except Exception:
            pass

    if not name:
        vals = [str(ws.cell(rows[0], i + 1).value or "").strip()
                for i in range(min(4, ws.max_column))]
        name = " ".join(v for v in vals if v and not v.isdigit())
        detail["full_name"] = name

    for r in sorted(rows, reverse=True):
        ws.delete_rows(r)
    try:
        wb.save(path)
    except PermissionError:
        return False, "users.xlsx Excelda ochiq — yopib qayta urining", detail
    detail["cleaned"]["users.xlsx"] = len(rows)

    # roles keshini yangilashga majburlaymiz
    try:
        import roles
        roles.reload()
    except Exception:
        pass

    # ── 2) Hisob va tokenlar ──
    try:
        import accounts
        if accounts.get_account(c):
            accounts.delete_account(c)
            detail["cleaned"]["accounts.json"] = 1
    except Exception as e:
        detail["cleaned"]["accounts.json"] = f"xato: {e}"

    # ── 3) Yordamchi fayllar ──
    for fn in DICT_FILES:
        p = BASE_DIR / fn
        d = _load(p)
        if isinstance(d, dict) and c in d:
            del d[c]
            _save(p, d)
            detail["cleaned"][fn] = 1
    for fn in RECORD_FILES:
        p = BASE_DIR / fn
        d = _load(p)
        if isinstance(d, dict) and isinstance(d.get("records"), list):
            before = len(d["records"])
            d["records"] = [x for x in d["records"]
                            if _norm(x.get("card_number")) != c]
            if len(d["records"]) != before:
                _save(p, d)
                detail["cleaned"][fn] = before - len(d["records"])

    # ── 4) Yashirilganlar ro'yxatidan ham olib tashlaymiz ──
    p = BASE_DIR / "dashboard_hidden.json"
    d = _load(p)
    if isinstance(d, dict) and isinstance(d.get("hidden_cards"), list):
        if c in [_norm(x) for x in d["hidden_cards"]]:
            d["hidden_cards"] = [x for x in d["hidden_cards"] if _norm(x) != c]
            _save(p, d)
            detail["cleaned"]["dashboard_hidden.json"] = 1

    # ── 5) Jonli kesh — dashboard darhol yangilansin ──
    p = BASE_DIR / CACHE_FILE
    d = _load(p)
    if isinstance(d, dict):
        n = 0
        for key, val in list(d.items()):
            if not isinstance(val, list):
                continue
            keep = [x for x in val
                    if not (isinstance(x, dict)
                            and _norm(x.get("card_number")) == c)]
            if len(keep) != len(val):
                n += len(val) - len(keep)
                d[key] = keep
        if n:
            s = d.get("summary")
            if isinstance(s, dict):
                for k in ("inside", "outside", "late", "absent",
                          "safari", "ruxsat", "tadbir", "interns"):
                    if k in d and isinstance(d[k], list):
                        s[k] = len(d[k])
            _save(p, d)
            detail["cleaned"][CACHE_FILE] = n

    return True, f"{name or c} tizimdan chiqarildi", detail


if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        print("Ishlatish: python xodim_chiqarish.py <JSHSHIR> [JSHSHIR2 ...]")
        raise SystemExit(1)
    for arg in sys.argv[1:]:
        ok, msg, det = remove_employee(arg)
        print(("OK   " if ok else "XATO ") + msg)
        for k, v in (det.get("cleaned") or {}).items():
            print(f"       {k}: {v}")
